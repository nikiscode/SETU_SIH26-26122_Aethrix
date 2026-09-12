"""Tests for the behaviour the demo actually rests on."""
from datetime import date

import pytest

from setu.accrual import accrue
from setu.linker import Linker
from setu.models import AUTO_COMMIT, PlanActivity, ProgressEvent, band_for
from setu.tags import (action_stems, extract_tags, normalise_text,
                       phrase_skeleton, section_of)


# --- tag grammar -----------------------------------------------------------

@pytest.mark.parametrize("text,expected", [
    ('Erect Line 24"-PG-1002 Spools', 'line:24"-PG-1002'),
    ("24 inch PG 1002 erected", 'line:24"-PG-1002'),
    ("poured FDN-C-0231", "foundation:FDN-C-0231"),
    ("cable laying done CBL-EL-0456", "cable:CBL-EL-0456"),
    ("loop checking completed LP-IN-0789", "loop:LP-IN-0789"),
    ("excavation near CH 12+300", "chainage:CH 12+300"),
])
def test_tag_extraction(text, expected):
    assert expected in extract_tags(text)


def test_line_tag_survives_notation_change():
    """The whole point: two notations for one line must canonicalise."""
    a = extract_tags('Erect Line 24"-PG-1002 Spools')
    b = extract_tags("welding done on 24 inch PG 1002")
    assert set(a) & set(b)


def test_normalise_collapses_verb_forms():
    assert "erect" in normalise_text("spool erection completed")
    assert "erect" in normalise_text("erected spools")
    assert "test" in normalise_text("hydro testing carried out")


def test_section_and_skeleton():
    assert section_of("welding completed Sec-3") == "sec-3"
    # the skeleton drops the volatile section number so a correction on
    # Sec-3 can transfer to Sec-4
    assert (phrase_skeleton("welding completed Sec-3")
            == phrase_skeleton("welding completed Sec-4"))


def test_action_stems_discriminate_work_kinds():
    assert action_stems("Hydrotest Line X") != action_stems("Erect Line X")


# --- linking ---------------------------------------------------------------

def _plan():
    return [
        PlanActivity("A-001", "P > Area 1 > Sec-3 > Piping",
                     'Erect Line 24"-PG-1002 Spools (Sec-3)', "Piping", "L6",
                     date(2026, 4, 1), date(2026, 4, 12), 100.0, "m",
                     ['line:24"-PG-1002']),
        PlanActivity("A-002", "P > Area 1 > Sec-3 > Piping",
                     'Hydrotest Line 24"-PG-1002 (Sec-3)', "Piping", "L6",
                     date(2026, 5, 1), date(2026, 5, 9), 100.0, "m",
                     ['line:24"-PG-1002']),
        PlanActivity("A-003", "P > Area 2 > Sec-3 > Piping",
                     'Erect Line 8"-CW-1200 Spools (Sec-3)', "Piping", "L6",
                     date(2026, 4, 1), date(2026, 4, 12), 100.0, "m",
                     ['line:8"-CW-1200']),
    ]


def _ev(text, **kw):
    kw.setdefault("event_date", date(2026, 4, 5))
    kw.setdefault("reported_on", kw["event_date"])
    return ProgressEvent("E1", text, "dpr_text", "t:1",
                         tags=extract_tags(text), **kw)


def test_tag_and_action_pick_the_right_activity():
    """Same line, same section — only the kind of work separates them."""
    lk = Linker(_plan())
    r = lk.link(_ev("spool erection completed on 24 inch PG 1002",
                    area="Area 1", discipline="Piping"))
    assert r.activity_id == "A-001"
    assert r.band == "auto_commit"


def test_action_verb_flips_the_answer():
    lk = Linker(_plan())
    r = lk.link(_ev("hydro testing carried out on 24 inch PG 1002",
                    area="Area 1", discipline="Piping",
                    event_date=date(2026, 5, 4)))
    assert r.activity_id == "A-002"


def test_area_disambiguates_identical_wording():
    """Identical text in two areas must resolve by the stated area."""
    lk = Linker(_plan())
    r = lk.link(_ev("erected spools at Sec-3", area="Area 2",
                    discipline="Piping"))
    assert r.activity_id == "A-003"


def test_out_of_plan_work_is_flagged_not_forced():
    lk = Linker(_plan())
    r = lk.link(_ev("temporary shoring installed at north access road"))
    assert r.band == "unmatched"
    assert r.activity_id is None


def test_confidence_bands_are_ordered():
    assert band_for(AUTO_COMMIT + 0.01) == "auto_commit"
    assert band_for(0.41) == "review"
    assert band_for(0.10) == "unmatched"


def test_alias_must_still_satisfy_the_schedule():
    """A confirmed phrasing does not license an impossible area."""
    lk = Linker(_plan())
    lk.learn("erected spools at Sec-3", "A-001")
    r = lk.link(_ev("erected spools at Sec-3", area="Area 2",
                    discipline="Piping"))
    assert r.activity_id != "A-001"


# --- accrual ---------------------------------------------------------------

def test_partial_entries_accrue_rather_than_compete():
    """Granularity mismatch: three field entries, one L6 node."""
    plan = _plan()
    lk = Linker(plan)
    events, links = [], {}
    for i, qty in enumerate([40.0, 40.0, 20.0], start=1):
        ev = ProgressEvent(f"E{i}", "spool erection on 24 inch PG 1002",
                           "dpr_text", f"t:{i}", date(2026, 4, 2 + i),
                           area="Area 1", discipline="Piping",
                           tags=extract_tags("24 inch PG 1002"),
                           quantity=qty, uom="m",
                           event_date=date(2026, 4, 2 + i),
                           action="completed" if qty == 20.0 else "progressed")
        events.append(ev)
        links[ev.event_id] = lk.link(ev)

    state = accrue(plan, events, links)
    st = state["A-001"]
    assert st.accrued_qty == pytest.approx(100.0)
    assert st.pct_complete == 1.0
    assert st.actual_start == date(2026, 4, 3)
    assert st.actual_finish == date(2026, 4, 5)
    assert len(st.event_ids) == 3


def test_review_band_is_not_credited_to_the_schedule():
    """The band exists to keep uncertain work out of the baseline."""
    plan = _plan()
    lk = Linker(plan)
    ev = ProgressEvent("E9", "some vague site work", "voice_chat", "v:1",
                       date(2026, 4, 4), quantity=10.0, uom="m",
                       event_date=date(2026, 4, 4))
    lr = lk.link(ev)
    lr.band = "review"
    lr.activity_id = "A-001"
    assert accrue(plan, [ev], {"E9": lr}) == {}
    assert accrue(plan, [ev], {"E9": lr}, include_review=True)["A-001"]


# --- reproducibility -------------------------------------------------------

def test_generated_inputs_are_byte_reproducible(tmp_path):
    """Same input, same output — including the sample files.

    Both writers had a source of per-process randomness that the seeded
    generator could not control: the voice log took its clock times from
    builtin hash(), which Python randomises per process, and the xlsx is a
    zip whose member dates and docProps/core.xml carry wall-clock time. The
    parsed events were identical either way, so no metric ever moved, but
    every regeneration rewrote a committed file.
    """
    from setu.formats import (parse_spreadsheet, parse_voice_log,
                              write_spreadsheet, write_voice_log)
    from setu.plan import save_plan
    from setu.synth import build_plan, generate_events

    plan, meta = build_plan()
    save_plan(plan, tmp_path / "plan.csv")
    df = generate_events(plan, meta)

    def written(name, writer):
        first = tmp_path / f"1_{name}"
        second = tmp_path / f"2_{name}"
        writer(df, first)
        writer(df, second)
        return first.read_bytes(), second.read_bytes()

    a, b = written("voice.txt", write_voice_log)
    assert a == b, "voice log differs between two runs of one generator"

    a, b = written("register.xlsx", write_spreadsheet)
    assert a == b, "xlsx differs between two runs of one generator"

    # and the files still parse to the same events they always did
    assert len(parse_voice_log(tmp_path / "1_voice.txt")) == 71
    assert len(parse_spreadsheet(tmp_path / "1_register.xlsx")) == 112


# --- robustness ------------------------------------------------------------

def test_noise_actually_degrades_and_is_seeded():
    """The noise model has to change the text, re-derive the tags from what
    it changed, and do the same thing twice for a given seed."""
    from setu.noise import perturb

    evs = [ProgressEvent(event_id="E1",
                         raw_text="terminations completed CBL-EL-0456",
                         source_format="dpr_text", source_ref="x:1",
                         reported_on=date(2026, 5, 1),
                         area="Gas Compression Area", discipline="Electrical",
                         tags=extract_tags("terminations completed "
                                           "CBL-EL-0456"),
                         quantity=40.0, uom="m",
                         event_date=date(2026, 5, 1))] * 40

    a, ta = perturb(evs, 1.0, seed=5)
    b, tb = perturb(evs, 1.0, seed=5)
    assert [e.raw_text for e in a] == [e.raw_text for e in b]
    assert ta == tb, "same seed must produce the same perturbations"

    assert ta["degraded"] > 0
    assert any(e.raw_text != evs[0].raw_text for e in a)

    # a tag that was dropped or mistyped must not survive in .tags, or the
    # noise would look harmless while the linker still saw the real tag
    for e in a:
        assert e.tags == extract_tags(e.raw_text)

    clean, tally = perturb(evs, 0.0, seed=5)
    assert [e.raw_text for e in clean] == [e.raw_text for e in evs]
    assert tally["degraded"] == 0


def test_noise_leaves_the_gold_label_alone():
    """A supervisor who mistypes a tag is still reporting the same work, so
    degradation must never silently rewrite what the right answer is."""
    from setu.noise import perturb
    from setu.formats import load_all
    from pathlib import Path

    evs = load_all(Path("data/inputs"))
    noisy, _ = perturb(evs, 0.8, seed=9)
    assert len(noisy) == len(evs)
    for a, b in zip(evs, noisy):
        assert a.event_id == b.event_id
        assert a.true_activity_id == b.true_activity_id


def test_template_split_holds_back_whole_phrasings():
    """The date split leaks — every phrasing appears on both sides of a
    chronological cut. A wording split must partition cleanly and hold back
    one phrasing per activity kind."""
    from setu.evaluate import load_templates, template_split
    from setu.pipeline import Pipeline
    from pathlib import Path

    templates = load_templates(Path("data/gold_labels.csv"))
    assert templates, "gold file carries no template ids"

    res = Pipeline(Path("data/baseline_schedule.csv")).run(Path("data/inputs"))
    seen, hold, held = template_split(res, templates)

    assert seen and hold
    assert not (seen & hold), "an event cannot be both seen and held out"
    assert seen | hold == set(res["links"]), "every event must be assigned"

    # exactly one phrasing held out per activity kind
    kinds = {t.rsplit(":", 1)[0] for t in templates.values() if t}
    assert len(held) == len(kinds)
    assert {t.rsplit(":", 1)[0] for t in held} == kinds

    # and no held-out phrasing produced any event on the seen side
    for eid in seen:
        assert templates.get(eid) not in held


# --- Primavera P6 -----------------------------------------------------------

def test_xer_round_trip_preserves_every_field(tmp_path):
    """The PS names Primavera exports as an input. The reader must recover
    everything the linker needs from a real XER, and the writer must
    produce one the reader can read back — byte-stably."""
    from pathlib import Path as P
    from setu.plan import load_plan, load_xer, save_xer
    plan = load_plan(P("data/baseline_schedule.csv"))
    x = tmp_path / "p.xer"
    save_xer(plan, x)
    back = load_xer(x)
    assert len(back) == len(plan) == 324
    for a, b in zip(plan, back):
        for f in ("activity_id", "wbs_path", "description", "discipline",
                  "level", "planned_start", "planned_finish", "planned_qty",
                  "uom", "tags", "predecessors"):
            assert getattr(a, f) == getattr(b, f), (a.activity_id, f)
    first = x.read_bytes()
    save_xer(plan, x)
    assert x.read_bytes() == first, "XER is not byte-deterministic"
    head = first.decode().splitlines()[0]
    assert head.startswith("ERMHDR\t")


def test_xer_reader_survives_a_minimal_real_export(tmp_path):
    """No UDFs at all — the way a plain P6 export arrives. Discipline must
    fall back to the WBS leaf and nothing may crash."""
    x = tmp_path / "plain.xer"
    x.write_text("\n".join([
        "ERMHDR\t19.12\t2026-01-01\tProject\tu\tu\tdb\tProject Management\tINR",
        "%T\tPROJECT", "%F\tproj_id\tproj_short_name", "%R\t1\tX",
        "%T\tPROJWBS", "%F\twbs_id\tproj_id\tparent_wbs_id\twbs_name\twbs_short_name",
        "%R\t10\t1\t\tOIL Project\tOIL", "%R\t11\t1\t10\tArea 9\tA9",
        "%R\t12\t1\t11\tPiping\tPIP",
        "%T\tTASK", "%F\ttask_id\tproj_id\twbs_id\ttask_code\ttask_name\t"
        "target_start_date\ttarget_end_date",
        "%R\t500\t1\t12\tX-001\tErect Line 6\"-CW-9001 Spools\t"
        "2026-05-01 08:00\t2026-05-09 17:00",
        "%E"]))
    from setu.plan import load_xer
    (a,) = load_xer(x)
    assert a.activity_id == "X-001"
    assert a.wbs_path == "OIL Project > Area 9 > Piping"
    assert a.discipline == "Piping"          # from the WBS leaf
    assert a.area == "Area 9"
    assert str(a.planned_start) == "2026-05-01"
    assert str(a.planned_finish) == "2026-05-09"


def test_p6_update_is_in_the_import_layout(tmp_path):
    """P6's spreadsheet import keys on internal field names in row 1 of a
    sheet called TASK. Get that wrong and the planner cannot import it."""
    from openpyxl import load_workbook
    from setu.plan import write_p6_update, P6_UPDATE_FIELDS
    rows = [
        dict(activity_id="A-1", description="x", actual_start="2026-05-01",
             actual_finish="2026-05-05", pct_complete=100.0),
        dict(activity_id="A-2", description="y", actual_start="",
             actual_finish="", pct_complete=0.0),        # nothing to say
    ]
    out = tmp_path / "u.xlsx"
    n = write_p6_update(rows, out)
    assert n == 1
    ws = load_workbook(out)["TASK"]
    got = list(ws.iter_rows(values_only=True))
    assert list(got[0]) == [f for f, _ in P6_UPDATE_FIELDS]
    assert list(got[1]) == [c for _, c in P6_UPDATE_FIELDS]
    assert got[2][0] == "A-1" and got[2][2] == "2026-05-01 08:00"
    assert len(got) == 3
    b = out.read_bytes(); write_p6_update(rows, out)
    assert out.read_bytes() == b, "update file is not byte-stable"


# --- the LLM stays out of the ranking path ----------------------------------

class _WrongButValidLLM:
    """Says the same valid area for everything — exactly what the 270M model
    did. A valid-but-wrong value is the case validation cannot catch."""
    model = "fake"
    def available(self): return True
    def extract(self, utterance, areas, disciplines):
        return {"area": areas[0], "discipline": disciplines[0],
                "action": "started"}


def test_llm_suggestions_never_reach_the_linker():
    """Measured: feeding model output to the linker as fact cost 14.3
    points of Top-1. So a suggestion pre-selects a chip and nothing more,
    and it must not suppress the question either."""
    from pathlib import Path as P
    from setu.agent import TimeAgent
    from setu.linker import Linker
    from setu.plan import load_plan
    plan = load_plan(P("data/baseline_schedule.csv"))
    linker = Linker(plan)
    plain = TimeAgent(plan, linker)
    with_llm = TimeAgent(plan, linker, llm=_WrongButValidLLM())

    utt = "cable termination done at Sec-1"
    a = plain.probe(plain.understand(utt))
    b_slots = with_llm.understand(utt)
    b = with_llm.probe(b_slots)

    # the model did fill slots — as suggestions
    assert "area" in b_slots.llm_filled and b_slots.parser == "rules+llm"
    # ...which the linker never saw: identical decision to rules-only
    assert b["link"].activity_id == a["link"].activity_id
    assert b["link"].confidence == a["link"].confidence
    # ...and the question is still asked, with the suggestion attached
    assert b["question"] == "area"
    assert b_slots.suggestion("area") == with_llm.areas[0]
    # the event it would commit carries no unconfirmed value
    assert with_llm.to_event(b_slots).area is None

    # once a person answers, the slot is confirmed and does count
    with_llm.answer(b_slots, "area", "Gas Compression Area")
    assert "area" not in b_slots.llm_filled
    assert with_llm.to_event(b_slots).area == "Gas Compression Area"


def test_llm_output_is_validated_against_the_plan():
    from setu.agent import LLMParser
    junk = {"area": "Mars", "discipline": "Alchemy", "quantity": "lots",
            "uom": "furlongs", "action": "flew", "extra": 1}
    assert LLMParser.validate(junk, ["Gas Compression Area"], ["Piping"]) == {}
    good = {"area": "Gas Compression Area", "discipline": "Piping",
            "quantity": 12, "uom": "m", "action": "completed"}
    assert LLMParser.validate(good, ["Gas Compression Area"], ["Piping"]) == \
        {"area": "Gas Compression Area", "discipline": "Piping",
         "quantity": 12.0, "uom": "m", "action": "completed"}
    assert LLMParser.validate("not a dict", [], []) == {}
    assert LLMParser.validate({"quantity": True}, [], []) == {}   # bool is not a qty


def test_activity_kind_strips_identifiers():
    from setu.memory import activity_kind
    from setu.models import PlanActivity
    from datetime import date as D
    mk = lambda d, disc: PlanActivity("X", "P > A > S > " + disc, d, disc,  # noqa: E731
                                      "L6", D(2026, 1, 1), D(2026, 1, 2), 1, "m")
    assert activity_kind(mk('Erect Line 24"-PG-1002 Spools (Sec-2)', "Piping")) \
        == "Piping · Erect Line Spools"
    assert activity_kind(mk('Erect Line 6"-CW-3301 Spools (North Bay)', "Piping")) \
        == "Piping · Erect Line Spools"
    assert activity_kind(mk("Pour Foundation FDN-C-0231 (Sec-1)", "Civil")) \
        == "Civil · Pour Foundation"
    assert activity_kind(mk("Loop Check LP-IN-0789 (Sec-3)", "Instrumentation")) \
        == "Instrumentation · Loop Check"
