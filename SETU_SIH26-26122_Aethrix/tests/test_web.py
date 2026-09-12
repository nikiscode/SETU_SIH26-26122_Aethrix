"""Tests for the server: the behaviour a planner's trust rests on.

Each test here corresponds to a promise the product makes — an entry is
persisted and linked, a planner's decision overrides the model and is
credited to the schedule, rejected work stays visible, and the numbers on
the console are derived from the entries rather than a cached rollup that
could drift.
"""
import importlib
from datetime import date

import pytest
from fastapi.testclient import TestClient


@pytest.fixture()
def client(tmp_path, monkeypatch):
    """A server on a throwaway database, so tests never touch demo state."""
    # point the module at a throwaway db BEFORE import, so module-level
    # setup never touches demo state or an unwritable deployment path
    monkeypatch.setenv("SETU_DB", str(tmp_path / "t.db"))
    from setu.web import app as appmod
    importlib.reload(appmod)
    appmod.LINKER.aliases, appmod.LINKER.vocab = {}, {}
    with TestClient(appmod.app) as c:
        c.app_module = appmod
        # Every request now carries an identity. The planner can report and
        # verify, which is what most of these tests exercise; role-specific
        # behaviour has its own tests below.
        c.cookies.set("setu_token", _token(appmod, "planner"))
        yield c


def _token(appmod, role: str) -> str:
    return next(u["token"] for u in appmod.STORE.all_users()
                if u["role"] == role)


def _post(c, text, **kw):
    return c.post("/api/entry", json={"text": text, **kw}).json()


def test_plan_is_anchored_around_today(client):
    """A real-time demo against a plan that finished months ago would
    penalise every entry for being outside its window."""
    plan = client.app_module.PLAN
    lo = min(a.planned_start for a in plan)
    hi = max(a.planned_finish for a in plan)
    assert lo <= date.today() <= hi


def test_entry_is_linked_and_persisted(client):
    r = _post(client, "loop checking completed LP-IN-0779",
              discipline="Instrumentation", quantity=2, uom="nos")
    assert r["band"] == "auto_commit"
    assert r["activity_id"].endswith("0079")
    assert r["latency_ms"] > 0
    assert client.get("/api/state").json()["totals"]["entries"] == 1


def test_untagged_entry_goes_to_review_with_a_reason(client):
    r = _post(client, "cable termination done at Sec-1", quantity=20, uom="m")
    assert r["band"] in ("review", "unmatched")
    assert r["notes"], "a queued item must say why it stopped"


def test_out_of_plan_work_is_flagged_not_forced(client):
    r = _post(client, "temporary shoring installed at north access road",
              quantity=3, uom="nos")
    assert r["band"] == "unmatched"
    assert r["activity_id"] is None
    # and it stays visible to the planner rather than vanishing
    q = client.get("/api/state").json()["queue"]
    assert any(x["id"] == r["id"] for x in q)


def test_planner_confirmation_credits_the_schedule(client):
    r = _post(client, "cable termination done at Sec-1", quantity=20, uom="m")
    before = client.get("/api/state").json()["totals"]["activities"]
    target = "OIL-A1-ELE-L6-0055"
    out = client.post(f"/api/queue/{r['id']}/confirm",
                      json={"activity_id": target, "learn": True}).json()
    assert out["activity_id"] == target
    assert out["confidence"] == 1.0      # a human looked at it
    # Not the literal word "planner": the trail has to name the person,
    # because these records end up as evidence in a payment dispute.
    users = {u["user_id"]: u for u in client.app_module.STORE.all_users()}
    assert out["resolved_by"] in users
    assert users[out["resolved_by"]]["role"] == "planner"
    assert out["verified_by"] == out["resolved_by"]
    assert out["stage"] == "verified"
    after = client.get("/api/state").json()
    assert after["totals"]["activities"] >= before
    assert all(x["id"] != r["id"] for x in after["queue"])


def test_confirmation_is_learned(client):
    r = _post(client, "cable termination done at Sec-1", quantity=20, uom="m")
    client.post(f"/api/queue/{r['id']}/confirm",
                json={"activity_id": "OIL-A1-ELE-L6-0055", "learn": True})
    assert client.get("/api/state").json()["learned"] >= 1


def test_reject_keeps_the_entry_visible_as_new_work(client):
    r = _post(client, "site canteen relocated to gate 2", quantity=1,
              uom="nos")
    client.post(f"/api/queue/{r['id']}/reject")
    q = client.get("/api/state").json()["queue"]
    row = next(x for x in q if x["id"] == r["id"])
    assert row["band"] == "unmatched"
    assert row["proposed"] is None


def test_totals_are_derived_from_entries_not_cached(client):
    for i in range(4):
        _post(client, f"loop checking completed LP-IN-0779 batch {i}",
              discipline="Instrumentation", quantity=1, uom="nos")
    t = client.get("/api/state").json()["totals"]
    assert t["entries"] == 4
    assert t["auto_commit"] + t["review"] + t["unmatched"] == 4


def test_upload_ingests_a_daily_report(client, tmp_path):
    f = tmp_path / "dpr_test.txt"
    f.write_text(
        "================\n"
        "DAILY PROGRESS REPORT — OIL Field Development Project\n"
        "Area: Gas Compression Area\n"
        "Date: 04-Apr-2026\n"
        "Prepared by: R. Gogoi\n"
        "================\n\n"
        "INSTRUMENTATION\n"
        "  - loop checking completed LP-IN-0779 (2 nos)   [Ref: E00001]\n",
        encoding="utf-8")
    with f.open("rb") as fh:
        r = client.post("/api/upload",
                        files={"file": ("dpr_test.txt", fh, "text/plain")})
    assert r.status_code == 200
    assert r.json()["ingested"] == 1
    assert client.get("/api/state").json()["totals"]["entries"] == 1


def test_activity_search_supports_planner_correction(client):
    hits = client.get("/api/activities", params={"q": "terminate cable"}).json()
    assert hits and all("Terminate Cable" in h["description"] for h in hits)


def test_reset_clears_state(client):
    _post(client, "loop checking completed LP-IN-0779",
          discipline="Instrumentation")
    client.post("/api/reset")
    assert client.get("/api/state").json()["totals"]["entries"] == 0


def test_accrual_accumulator_matches_a_full_recompute(client):
    """The incremental accrual cache must equal folding the whole table.

    That cache exists because re-folding every stored entry on each ingest
    made per-entry latency grow with the table (4.0 ms at 335 rows, 24.7 ms
    at 2,677 — quadratic overall). It is the one derived value this module
    keeps, so it has to be provably identical to the thing it replaces or
    the remaining-quantity constraint starts lying.
    """
    appmod = client.app_module
    for text in ('spools erected on line 12"-CW-1001',
                 "poured FDN-C-0231", "terminated both ends CBL-EL-0457",
                 'spools erected on line 12"-CW-1001'):
        _post(client, text, quantity=10, uom="m")

    incremental = dict(appmod.ACCRUED)
    appmod._rebuild_accrued()
    assert incremental == appmod.ACCRUED
    assert incremental, "nothing accrued — the test proves nothing"


def test_planner_resolve_keeps_the_accumulator_honest(client):
    """A confirm moves an entry into the credited set, so the accumulator
    must follow — otherwise the quantity constraint drifts as planners work."""
    appmod = client.app_module
    r = _post(client, "cable termination done at Sec-1", quantity=12, uom="m")
    assert r["band"] == "review"          # not credited yet
    client.post(f"/api/queue/{r['id']}/confirm",
                json={"activity_id": r["candidates"][0], "learn": False})

    incremental = dict(appmod.ACCRUED)
    appmod._rebuild_accrued()
    assert incremental == appmod.ACCRUED


def test_reset_requires_the_token_when_one_is_configured(tmp_path,
                                                         monkeypatch):
    """Reset deletes every entry. Bound to anything but loopback that is a
    one-request wipe of a live demo, so the launcher configures a token and
    this endpoint has to honour it."""
    monkeypatch.setenv("SETU_DB", str(tmp_path / "guard.db"))
    monkeypatch.setenv("SETU_ADMIN_TOKEN", "topsecret")
    from setu.web import app as appmod
    importlib.reload(appmod)
    with TestClient(appmod.app) as c:
        c.cookies.set("setu_token", _token(appmod, "planner"))
        _post(c, 'spools erected on line 12"-CW-1001')
        assert c.get("/api/state").json()["totals"]["entries"] == 1

        assert c.post("/api/reset").status_code == 403
        assert c.post("/api/reset",
                      headers={"X-SETU-Token": "wrong"}).status_code == 403
        assert c.get("/api/state").json()["totals"]["entries"] == 1

        assert c.post("/api/reset",
                      headers={"X-SETU-Token": "topsecret"}).status_code == 200
        assert c.get("/api/state").json()["totals"]["entries"] == 0


# --- roles and the approval chain -----------------------------------------

def _as(client, role):
    client.cookies.set("setu_token", _token(client.app_module, role))
    return client


def test_no_identity_no_write(client):
    """Every mutating endpoint needs a signed-in person. An audit trail is
    worthless if anyone on the network can post progress as anyone."""
    client.cookies.clear()
    assert client.post("/api/entry", json={"text": "poured FDN-C-0231"}
                       ).status_code == 401
    assert client.post("/api/queue/1/confirm", json={}).status_code == 401
    assert client.post("/api/approve", json={"entry_ids": [1]}
                       ).status_code == 401


def test_roles_cannot_do_each_others_jobs(client):
    """Capability is checked on the server. A supervisor who finds the
    approve endpoint must not be able to approve their own work."""
    r = _as(client, "supervisor").post(
        "/api/entry", json={"text": "poured FDN-C-0231", "quantity": 6,
                            "uom": "m3"})
    assert r.status_code == 200
    eid = r.json()["id"]

    # supervisor: report only
    assert client.post("/api/queue/%d/confirm" % eid,
                       json={}).status_code == 403
    assert client.post("/api/approve",
                       json={"entry_ids": [eid]}).status_code == 403
    assert client.post("/api/certify",
                       json={"entry_ids": [eid]}).status_code == 403

    # planner verifies but cannot approve or certify
    _as(client, "planner")
    assert client.post("/api/approve",
                       json={"entry_ids": [eid]}).status_code == 403

    # owner certifies but must not be able to verify a link
    _as(client, "owner")
    assert client.post("/api/queue/%d/confirm" % eid,
                       json={}).status_code == 403

    # auditor changes nothing at all
    _as(client, "auditor")
    assert client.post("/api/entry",
                       json={"text": "x"}).status_code == 403
    assert client.post("/api/approve",
                       json={"entry_ids": [eid]}).status_code == 403


def test_the_chain_cannot_be_skipped(client):
    """Progress must not reach a payment certificate without passing
    through verification and approval."""
    eid = _as(client, "supervisor").post(
        "/api/entry", json={"text": "cable termination done at Sec-1",
                            "quantity": 12, "uom": "m"}).json()["id"]

    # straight to certify: refused, because nothing has verified it
    assert _as(client, "owner").post(
        "/api/certify", json={"entry_ids": [eid]}).status_code == 409
    # approve before verify: refused too
    assert _as(client, "pm").post(
        "/api/approve", json={"entry_ids": [eid]}).status_code == 409

    # walk it properly
    r = _as(client, "planner").post("/api/queue/%d/confirm" % eid, json={})
    assert r.status_code == 200 and r.json()["stage"] == "verified"
    assert _as(client, "pm").post(
        "/api/approve", json={"entry_ids": [eid]}).status_code == 200
    out = _as(client, "owner").post("/api/certify",
                                    json={"entry_ids": [eid]})
    assert out.status_code == 200 and out.json()["certified"] == [eid]


def test_audit_reconstructs_the_whole_chain(client):
    """The auditor's question is 'who turned this sentence into a
    certified quantity, and when'. Every stage must answer with a name."""
    eid = _as(client, "supervisor").post(
        "/api/entry", json={"text": "loop checking completed LP-IN-0779",
                            "quantity": 2, "uom": "nos"}).json()["id"]
    _as(client, "planner").post("/api/queue/%d/confirm" % eid, json={})
    _as(client, "pm").post("/api/approve", json={"entry_ids": [eid]})
    _as(client, "owner").post("/api/certify", json={"entry_ids": [eid]})

    a = _as(client, "auditor").get("/api/audit/entry/%d" % eid).json()
    stages = {c["stage"]: c for c in a["chain"]}
    assert {"reported", "verified", "approved", "certified"} <= set(stages)
    for st in ("verified", "approved", "certified"):
        assert stages[st]["name"], f"{st} has no attributed person"
        assert stages[st]["at"], f"{st} has no timestamp"
    # and the action log corroborates it independently of entry columns
    actions = [l["action"] for l in a["log"]]
    assert ["report", "verify", "approve", "certify"] == actions
    assert all(l["user_name"] for l in a["log"])


def test_public_role_sees_no_commercial_detail(client):
    """Transparency must not become a commercial leak, or expose one
    person's output."""
    from setu.roles import ROLES
    pub = ROLES["public"]
    assert pub.aggregate_only and pub.read_only
    assert not pub.sees_quantities and not pub.sees_raw_text

    from setu.roles import User
    u = User("p", "Public", "public")
    kept = u.redact({"raw_text": "x", "quantity": 5, "uom": "m",
                     "confidence": 0.9, "activity_id": "A"})
    assert "activity_id" in kept
    for gone in ("raw_text", "quantity", "uom", "confidence"):
        assert gone not in kept


def test_contractor_scope_excludes_other_organisations(client):
    from setu.roles import User
    a = User("c", "S. Baruah", "contractor", org="Contractor A")
    assert a.may_see_entry({"org": "Contractor A"})
    assert not a.may_see_entry({"org": "Contractor B"})
    # No org recorded means it might be a competitor's, so a scoped role
    # does not see it. It is not lost: every all-orgs role still does.
    assert not a.may_see_entry({"org": ""})
    from setu.roles import ROLES
    assert ROLES["planner"].sees_all_orgs and ROLES["auditor"].sees_all_orgs
    # area, unlike org, is a filter rather than a boundary
    b = User("d", "Foreman", "supervisor", org="Contractor A",
             areas=["Gas Compression Area"])
    assert b.may_see_entry({"org": "Contractor A", "area": ""})
    assert not b.may_see_entry({"org": "Contractor A",
                                "area": "Utilities & Flare"})


def test_agent_asks_nothing_when_the_tag_decides_it(client):
    """The difference from a form: it only asks when asking would help."""
    _as(client, "supervisor")
    r = client.post("/api/agent/turn",
                    json={"utterance": "loop checking completed LP-IN-0779",
                          "slots": {}}).json()
    assert r["proposal"]["band"] == "auto_commit"
    assert "question" not in r

    r2 = client.post("/api/agent/turn",
                     json={"utterance": "cable termination done at Sec-1",
                           "slots": {}}).json()
    assert r2["question"]["slot"] == "area"
    assert r2["question"]["options"], "a question with no answers to tap"


def test_agent_answer_improves_the_link_then_commits(client):
    _as(client, "supervisor")
    t1 = client.post("/api/agent/turn",
                     json={"utterance": "cable termination done at Sec-1",
                           "slots": {}}).json()
    area = t1["question"]["options"][0]
    t2 = client.post("/api/agent/turn",
                     json={"utterance": area, "slots": t1["slots"],
                           "answering": "area"}).json()
    assert t2["slots"]["area"] == area
    # The answer must actually constrain the choice: whatever it now
    # proposes has to sit in the area the supervisor just named. (Note it
    # does not have to be *more* confident — answering with a wrong area
    # should lower confidence, and does.)
    plan = {a.activity_id: a for a in client.app_module.PLAN}
    assert plan[t2["proposal"]["activity_id"]].area == area
    # and it does not ask the same question twice
    assert t2.get("question", {}).get("slot") != "area"

    row = client.post("/api/agent/commit",
                      json={"slots": t2["slots"]}).json()
    assert row["source"] == "time_agent"
    # Attribution comes from the session, never from the request body: a
    # client must not be able to file progress under someone else's name.
    me = client.get("/api/me").json()
    assert row["reporter"] == me["name"]
    assert row["reported_by"] == me["user_id"]
    assert row["org"] == me["org"]


def test_public_and_scoped_roles_cannot_read_the_audit_trail(client):
    """The public portal reading the action log would expose raw field text
    and one person's output over time — the exact harm the role exists to
    prevent. It was reachable by anyone signed in."""
    eid = _as(client, "supervisor").post(
        "/api/entry", json={"text": "poured FDN-C-0231", "quantity": 6,
                            "uom": "m3"}).json()["id"]

    for role in ("public", "contractor"):
        _as(client, role)
        assert client.get("/api/audit/log").status_code == 403

    # ...but a supervisor may still trace their own entry
    _as(client, "supervisor")
    assert client.get(f"/api/audit/entry/{eid}").status_code == 200
    assert client.get("/api/audit/log").status_code == 403

    for role in ("planner", "pm", "owner", "auditor", "head"):
        _as(client, role)
        assert client.get("/api/audit/log").status_code == 200, role


def test_cannot_certify_progress_with_no_measured_quantity(client):
    """Certification is what turns progress into money owed, so it needs a
    number. A certified entry with no quantity is a record worth nothing."""
    eid = _as(client, "supervisor").post(
        "/api/entry", json={"text": "cable termination done at Sec-1"}
    ).json()["id"]
    _as(client, "planner").post(f"/api/queue/{eid}/confirm", json={})
    _as(client, "pm").post("/api/approve", json={"entry_ids": [eid]})
    r = _as(client, "owner").post("/api/certify", json={"entry_ids": [eid]})
    assert r.status_code == 422
    assert "quantity" in r.json()["detail"]

    # with a quantity it certifies
    eid2 = _as(client, "supervisor").post(
        "/api/entry", json={"text": "loop checking completed LP-IN-0779",
                            "quantity": 2, "uom": "nos"}).json()["id"]
    _as(client, "planner").post(f"/api/queue/{eid2}/confirm", json={})
    _as(client, "pm").post("/api/approve", json={"entry_ids": [eid2]})
    ok = _as(client, "owner").post("/api/certify", json={"entry_ids": [eid2]})
    assert ok.status_code == 200 and ok.json()["quantity"] == 2.0


def test_upload_requires_identity_too(client):
    """Every write path is gated, or the gates on the others are theatre."""
    import io
    doc = io.BytesIO(b"""================
DAILY PROGRESS REPORT
Area: Gas Compression Area
Date: 01-Jun-2026
Prepared by: R. Gogoi
================

PIPING
  - spools erected on line 12"-CW-1001 (20 m)   [Ref: E90001]
""")
    client.cookies.clear()
    r = client.post("/api/upload",
                    files={"file": ("dpr_x.txt", doc, "text/plain")})
    assert r.status_code == 401

    doc.seek(0)
    _as(client, "supervisor")
    r = client.post("/api/upload",
                    files={"file": ("dpr_x.txt", doc, "text/plain")})
    assert r.status_code == 200 and r.json()["ingested"] == 1
    # and the ingested rows are attributed to the uploader
    me = client.get("/api/me").json()
    rows = client.get("/api/state").json()["recent"]
    assert rows, "nothing ingested"
    e = client.app_module.STORE.entry(rows[0]["id"])
    assert e["reported_by"] == me["user_id"]


# --- institutional memory ---------------------------------------------------

def test_memory_is_derived_and_gated(client):
    """PS outcome 5(b): a queryable record of what actually happened. It is
    computed from the entries, and it exposes discipline productivity, so
    the public and scoped roles do not get it."""
    _as(client, "supervisor")
    for text, q in [('spools erected on line 12"-CW-1001', 30),
                    ('spools erected on line 12"-CW-1001', 30)]:
        client.post("/api/entry", json={"text": text, "quantity": q,
                                        "uom": "m"})
    for role in ("public", "contractor", "supervisor"):
        _as(client, role)
        assert client.get("/api/memory").status_code == 403, role

    m = _as(client, "pm").get("/api/memory").json()
    assert m["scope"]["entries"] == 2
    assert m["scope"]["with_actuals"] >= 1
    assert {"durations", "productivity", "slippage_by_area",
            "slippage_by_discipline", "causes", "cause_taxonomy"} <= set(m)
    assert any(t["key"] == "material" for t in m["cause_taxonomy"])


def test_delay_cause_attached_at_verification_reaches_memory(client):
    """The planner's 'why' is recorded once, at verification, and shows up
    aggregated — instead of dying as a remark in a DPR column."""
    eid = _as(client, "supervisor").post(
        "/api/entry", json={"text": "cable termination done at Sec-1",
                            "quantity": 12, "uom": "m"}).json()["id"]
    r = _as(client, "planner").post(
        f"/api/queue/{eid}/confirm",
        json={"delay_cause": "drawing"})
    assert r.status_code == 200 and r.json()["delay_cause"] == "drawing"

    # junk causes are dropped, not stored
    eid2 = _as(client, "supervisor").post(
        "/api/entry", json={"text": "cable termination done at Sec-2",
                            "quantity": 5, "uom": "m"}).json()["id"]
    r2 = _as(client, "planner").post(f"/api/queue/{eid2}/confirm",
                                     json={"delay_cause": "aliens"})
    assert r2.status_code == 200 and not r2.json().get("delay_cause")

    m = _as(client, "head").get("/api/memory").json()
    assert m["causes"]["recorded"] == 1
    assert m["causes"]["ranked"][0]["cause"] == "drawing"
    # and it is in the audit detail too
    log = _as(client, "auditor").get(f"/api/audit/entry/{eid}").json()["log"]
    assert any("cause=drawing" in (l["detail"] or "") for l in log)


def test_p6_writeback_carries_only_signed_off_progress(client):
    """The file a planner drags into Primavera must contain what the chain
    approved — not what the model merely guessed."""
    from openpyxl import load_workbook
    import io
    _as(client, "supervisor")
    a = client.post("/api/entry", json={"text": "loop checking completed LP-IN-0779",
                                        "quantity": 2, "uom": "nos"}).json()
    b = client.post("/api/entry", json={"text": "poured FDN-C-0231",
                                        "quantity": 6, "uom": "m3"}).json()
    assert a["band"] == "auto_commit" and b["band"] == "auto_commit"
    # only `a` is walked through verify + approve
    _as(client, "planner").post(f"/api/queue/{a['id']}/confirm", json={})
    _as(client, "pm").post("/api/approve", json={"entry_ids": [a["id"]]})

    assert _as(client, "planner").get("/api/writeback/p6.xlsx").status_code == 403
    r = _as(client, "pm").get("/api/writeback/p6.xlsx")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content))["TASK"]
    ids = [row[0] for row in ws.iter_rows(min_row=3, values_only=True)]
    assert a["activity_id"] in ids
    assert b["activity_id"] not in ids, "unapproved progress leaked into P6"
